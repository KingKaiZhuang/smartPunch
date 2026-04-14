"""
Excel 匯出相關功能 - 支援多種格式（改為使用 Excel 模板，不再自行生成空白 Workbook）
"""
from flask import request, make_response
from openpyxl import load_workbook
from datetime import datetime
import io
import os
from db import get_connection
from flask_login import login_required
from collections import defaultdict


# ========= 模板設定 =========
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ✅ 請把你轉好的模板放在：<本檔案同層>/templates/底下
TEMPLATE_PATH = os.path.join(BASE_DIR, "templates", "志工時數匯入.xlsx")  # 同一份模板同時含「明細」與「統計」工作表

# ✅ 若你的工作表名稱不同，改這裡
SHEET_DETAILED_NAME = "服務時數記錄"
SHEET_SUMMARY_NAME = "服務統計"


def _load_template(template_path: str, preferred_sheet=None):
    """載入模板並回傳 (wb, ws)。若 preferred_sheet 不存在就用 active。"""
    if not os.path.exists(template_path):
        return None, None

    wb = load_workbook(template_path)
    if preferred_sheet and preferred_sheet in wb.sheetnames:
        ws = wb[preferred_sheet]
    else:
        ws = wb.active
    return wb, ws


def _clear_values_keep_style(ws, start_row: int = 2):
    """清空 start_row 之後的儲存格 value（保留樣式、欄寬、框線、凍結窗格等模板設定）。"""
    max_row = ws.max_row
    if max_row >= start_row:
        for row in ws.iter_rows(min_row=start_row, max_row=max_row):
            for cell in row:
                cell.value = None


def _build_filename(nid: str, date_start, date_end, tag: str, ext: str = "xlsx") -> str:
    filename_parts = []
    if nid:
        filename_parts.append(nid)

    if date_start and date_end:
        filename_parts.append(f"{date_start}_to_{date_end}")
    elif date_start:
        filename_parts.append(f"from_{date_start}")
    elif date_end:
        filename_parts.append(f"to_{date_end}")

    filename_parts.append(f"{tag}_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
    return "_".join(filename_parts) + f".{ext}"


def register_export_routes(app):
    """註冊匯出相關的路由"""

    # ----------- 匯出資料（支援多格式）-----------
    @app.route('/export_xlsx', methods=['GET', 'POST'])
    @login_required
    def export_xlsx():
        nid = request.values.get('nid', '').strip()
        date_start = request.values.get('date_start')
        date_end = request.values.get('date_end')
        export_format = request.values.get('format', 'detailed')  # detailed / summary

        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        query = """
            SELECT name, id_number, service_start, service_end,
                service_item, service_content, service_hours, service_minutes,
                served_people_count, transport_fee, meal_fee, service_area,
                remarks, import_action, serial_number, foreign_service_count, domestic_service_count
            FROM service_records
            WHERE 1=1
        """
        params = []

        if nid:
            query += " AND LOWER(id_number) = LOWER(%s)"
            params.append(nid)

        if date_start and date_end:
            query += " AND DATE(service_start) BETWEEN %s AND %s"
            params.extend([date_start, date_end])
        elif date_start:
            query += " AND DATE(service_start) >= %s"
            params.append(date_start)
        elif date_end:
            query += " AND DATE(service_start) <= %s"
            params.append(date_end)

        query += " ORDER BY service_start DESC"

        cursor.execute(query, tuple(params))
        records = cursor.fetchall()
        cursor.close()
        conn.close()

        if export_format == 'summary':
            return export_summary_format(records, nid, date_start, date_end)
        else:
            return export_detailed_format(records, nid, date_start, date_end)


def export_detailed_format(records, nid, date_start, date_end):
    """詳細格式（逐筆記錄）- ✅ 使用模板，不自行生成 Workbook"""

    wb, ws = _load_template(TEMPLATE_PATH, SHEET_DETAILED_NAME)
    if wb is None:
        return make_response(
            f"找不到詳細模板檔案：{TEMPLATE_PATH}\n請先把 .xls 另存為 .xlsx，並放到 templates/ 底下。",
            500
        )

    # 清掉舊資料（保留模板樣式）
    _clear_values_keep_style(ws, start_row=2)

    # 模板欄位順序 A~Q（共 17 欄）
    # A 姓名
    # B 身分證字號
    # C 服務日期起（service_start）
    # D 服務日期迄（service_end）
    # E 服務項目
    # F 服務內容
    # G 服務時數_小時
    # H 服務時數_分鐘
    # I 受服務人次
    # J 交通費
    # K 誤餐費
    # L 服務區域
    # M 備註
    # N 匯入動作
    # O 序號（serial_number）
    # P 國外參與服務人次
    # Q 國內參與服務人次

    start_row = 2
    if not records:
        ws.cell(row=start_row, column=1, value="無資料")
    else:
        for i, r in enumerate(records, start=start_row):
            ws.cell(row=i, column=1,  value=r.get("name") or "")
            ws.cell(row=i, column=2,  value=r.get("id_number") or "")
            ws.cell(row=i, column=3,  value=r.get("service_start") or "")
            ws.cell(row=i, column=4,  value=r.get("service_end") or "")
            ws.cell(row=i, column=5,  value=r.get("service_item") or "")
            ws.cell(row=i, column=6,  value=r.get("service_content") or "")
            ws.cell(row=i, column=7,  value=r.get("service_hours") or 0)
            ws.cell(row=i, column=8,  value=r.get("service_minutes") or 0)
            ws.cell(row=i, column=9,  value=r.get("served_people_count") or 0)
            ws.cell(row=i, column=10, value=r.get("transport_fee") or 0)
            ws.cell(row=i, column=11, value=r.get("meal_fee") or 0)
            ws.cell(row=i, column=12, value=r.get("service_area") or "")
            ws.cell(row=i, column=13, value=r.get("remarks") or "")
            ws.cell(row=i, column=14, value=r.get("import_action") or "")
            ws.cell(row=i, column=15, value=r.get("serial_number") or "")
            ws.cell(row=i, column=16, value=r.get("foreign_service_count") or 0)
            ws.cell(row=i, column=17, value=r.get("domestic_service_count") or 0)

    filename = _build_filename(nid, date_start, date_end, tag="detailed", ext="xlsx")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response


def export_summary_format(records, nid, date_start, date_end):
    """統計格式（按人員統計）- ✅ 使用模板，不自行生成 Workbook"""

    wb, ws = _load_template(TEMPLATE_PATH, SHEET_SUMMARY_NAME)
    if wb is None:
        return make_response(
            f"找不到統計模板檔案：{TEMPLATE_PATH}\n請先準備統計用的 .xlsx 模板並放到 templates/ 底下。",
            500
        )

    # 按姓名分組並統計
    grouped_data = defaultdict(lambda: {
        'id_number': '',
        'service_dates': [],
        'service_items': set(),
        'service_contents': set(),
        'total_hours': 0,
        'total_minutes': 0,
        'total_people': 0,
        'total_transport': 0,
        'total_meal': 0,
        'service_areas': set(),
        'remarks': [],
        'import_actions': set(),
        'serials': [],
        'foreign_count': 0,
        'domestic_count': 0
    })

    for r in records:
        name = r.get('name') or ""
        if not name:
            continue

        grouped_data[name]['id_number'] = r.get('id_number') or ""

        if r.get('service_start'):
            try:
                service_date = r['service_start'].strftime('%Y-%m-%d')
            except Exception:
                service_date = str(r.get('service_start'))
            grouped_data[name]['service_dates'].append(service_date)

        if r.get('service_item'):
            grouped_data[name]['service_items'].add(str(r['service_item']))
        if r.get('service_content'):
            grouped_data[name]['service_contents'].add(str(r['service_content']))

        grouped_data[name]['total_hours'] += (r.get('service_hours') or 0)
        grouped_data[name]['total_minutes'] += (r.get('service_minutes') or 0)
        grouped_data[name]['total_people'] += (r.get('served_people_count') or 0)
        grouped_data[name]['total_transport'] += (r.get('transport_fee') or 0)
        grouped_data[name]['total_meal'] += (r.get('meal_fee') or 0)

        if r.get('service_area'):
            grouped_data[name]['service_areas'].add(str(r['service_area']))

        if r.get('remarks'):
            grouped_data[name]['remarks'].append(str(r['remarks']))

        if r.get('import_action'):
            grouped_data[name]['import_actions'].add(str(r['import_action']))

        serial_num = r.get('serial_number') or ""
        if serial_num:
            grouped_data[name]['serials'].append(str(serial_num))

        grouped_data[name]['foreign_count'] += (r.get('foreign_service_count') or 0)
        grouped_data[name]['domestic_count'] += (r.get('domestic_service_count') or 0)

    # 處理分鐘進位
    for name in grouped_data:
        data = grouped_data[name]
        if data['total_minutes'] >= 60:
            extra_hours = data['total_minutes'] // 60
            data['total_hours'] += extra_hours
            data['total_minutes'] = data['total_minutes'] % 60

    # 清掉舊資料（保留模板樣式）
    _clear_values_keep_style(ws, start_row=2)

    start_row = 2
    if not grouped_data:
        ws.cell(row=start_row, column=1, value="無資料")
    else:
        row_i = start_row
        for name in sorted(grouped_data.keys()):
            data = grouped_data[name]

            dates = sorted(data['service_dates'])
            date_start_str = dates[0] if dates else ""
            date_end_str = dates[-1] if dates else ""

            row = [
                name,
                data['id_number'],
                date_start_str,
                date_end_str,
                '、'.join(sorted(data['service_items'])) if data['service_items'] else "",
                '、'.join(sorted(data['service_contents'])) if data['service_contents'] else "",
                data['total_hours'],
                data['total_minutes'],
                data['total_people'],
                data['total_transport'],
                data['total_meal'],
                '、'.join(sorted(data['service_areas'])) if data['service_areas'] else "",
                '；'.join(data['remarks']) if data['remarks'] else "",
                '、'.join(sorted(data['import_actions'])) if data['import_actions'] else "",
                '、'.join(data['serials']) if data['serials'] else "",
                data['foreign_count'],
                data['domestic_count'],
            ]

            for col_i, val in enumerate(row, start=1):
                ws.cell(row=row_i, column=col_i, value=val)

            row_i += 1

    filename = _build_filename(nid, date_start, date_end, tag="summary", ext="xlsx")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers['Content-Disposition'] = f'attachment; filename="{filename}"'
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return response
