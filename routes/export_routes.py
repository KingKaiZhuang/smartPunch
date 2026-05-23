"""
Excel 匯出相關功能 - 支援多種格式（改為使用 Excel 模板，不再自行生成空白 Workbook）
"""
from flask import request, make_response
from openpyxl import load_workbook
from datetime import datetime
import io
import os
from db import get_connection
from flask_login import login_required
from collections import defaultdict
import subprocess
import tempfile
from datetime import datetime, date


# ========= 模板設定 =========
BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# ✅ 讀取專案根目錄下 templates 資料夾中的模板
PROJECT_ROOT = os.path.dirname(BASE_DIR)
TEMPLATE_PATH = os.path.join(PROJECT_ROOT, "templates", "志工時數匯入.xlsx")

# ✅ 若你的工作表名稱不同，改這裡
SHEET_DETAILED_NAME = "服務時數記錄"
SHEET_SUMMARY_NAME = "服務統計"

def minutes_to_0_or_30(hours, minutes):
    """
    時間進位邏輯：
    - 0-15 分鐘：0 分鐘
    - 16-30 分鐘：30 分鐘
    - 31-59 分鐘：進位 1 小時，0 分鐘
    """
    h = int(hours or 0)
    m = int(minutes or 0)

    # 換算成總分鐘數
    total_minutes = h * 60 + m
    
    # 算出基礎的「小時」與「剩餘的分鐘」
    base_h = total_minutes // 60
    rem_m = total_minutes % 60

    # 依照新規則進行區間判斷
    if rem_m <= 15:
        out_h = base_h
        out_m = 0
    elif rem_m <= 30:
        out_h = base_h
        out_m = 30
    else:
        # 31到59分鐘：小時數加 1，分鐘數歸零
        out_h = base_h + 1
        out_m = 0

    return out_h, out_m


def to_roc_yyyMMdd(v):
    """
    把日期/時間轉成民國 yyyMMdd（無-），例如 2026-01-01 -> 1150101
    會回傳字串，避免 Excel 亂轉格式
    """
    if v is None or v == "":
        return ""

    # v 可能是 datetime/date/或字串
    if isinstance(v, datetime):
        d = v.date()
    elif isinstance(v, date):
        d = v
    elif isinstance(v, str):
        s = v.strip()
        if not s:
            return ""
        # 只取日期部分，容忍 "YYYY-MM-DD HH:MM:SS"
        s = s.split(" ", 1)[0].replace("/", "-")
        try:
            d = datetime.strptime(s, "%Y-%m-%d").date()
        except:
            return v  # 解析失敗就原樣回傳
    else:
        return str(v)

    roc_year = d.year - 1911
    # 建議固定 3 碼年（台灣常用），1150101、0990101
    return f"{roc_year:03d}{d.month:02d}{d.day:02d}"

def _load_template(template_path: str, preferred_sheet=None):
    """載入模板並回傳 (wb, ws)。若 preferred_sheet 不存在就用 active。"""
    if not os.path.exists(template_path):
        return None, None

    wb = load_workbook(template_path)
    if preferred_sheet and preferred_sheet in wb.sheetnames:
        ws = wb[preferred_sheet]
    else:
        ws = wb.active
    return wb, ws


def _clear_values_keep_style(ws, start_row: int = 2):
    """清空 start_row 之後的儲存格 value（保留樣式、欄寬、框線、凍結窗格等模板設定）。"""
    max_row = ws.max_row
    if max_row >= start_row:
        for row in ws.iter_rows(min_row=start_row, max_row=max_row):
            for cell in row:
                cell.value = None


def _build_filename(name: str, nid: str, date_start, date_end, tag: str, ext: str = "xlsx") -> str:
    filename_parts = []
    if name:
        filename_parts.append(name)
    if nid:
        filename_parts.append(nid)

    if date_start and date_end:
        filename_parts.append(f"{date_start}_to_{date_end}")
    elif date_start:
        filename_parts.append(f"from_{date_start}")
    elif date_end:
        filename_parts.append(f"to_{date_end}")

    filename_parts.append(f"{tag}_{datetime.now().strftime('%Y%m%d_%H%M%S')}")
    return "_".join(filename_parts) + f".{ext}"


def register_export_routes(app):
    """註冊匯出相關的路由"""

    # ----------- 匯出資料（支援多格式）-----------
    @app.route('/export_xlsx', methods=['GET', 'POST'])
    @login_required
    def export_xlsx():
        name = request.values.get('name', '').strip()
        nid = request.values.get('nid', '').strip()
        date_start = request.values.get('date_start')
        date_end = request.values.get('date_end')
        export_format = request.values.get('format', 'detailed')  # detailed / summary

        conn = get_connection()
        cursor = conn.cursor(dictionary=True)

        query = """
            SELECT name, id_number, service_start, service_end,
                service_item, service_content, service_hours, service_minutes,
                served_people_count, transport_fee, meal_fee, service_area,
                remarks, import_action, serial_number, foreign_service_count, domestic_service_count
            FROM service_records
            WHERE 1=1
        """
        params = []

        if name:
            query += " AND name LIKE ?"
            params.append(f"%{name}%")

        if nid:
            query += " AND LOWER(id_number) = LOWER(?)"
            params.append(nid)

        if date_start and date_end:
            query += " AND DATE(service_start) BETWEEN ? AND ?"
            params.extend([date_start, date_end])
        elif date_start:
            query += " AND DATE(service_start) >= ?"
            params.append(date_start)
        elif date_end:
            query += " AND DATE(service_start) <= ?"
            params.append(date_end)

        query += " ORDER BY service_start DESC"

        cursor.execute(query, tuple(params))
        records = cursor.fetchall()
        cursor.close()
        conn.close()

        if export_format == 'summary':
            return export_summary_format(records, name, nid, date_start, date_end)
        else:
            return export_detailed_format(records, name, nid, date_start, date_end)


def export_detailed_format(records, name, nid, date_start, date_end):
    """詳細格式（逐筆記錄）- ✅ 使用模板，不自行生成 Workbook"""

    wb, ws = _load_template(TEMPLATE_PATH, SHEET_DETAILED_NAME)
    if wb is None:
        return make_response(
            f"找不到詳細模板檔案：{TEMPLATE_PATH}\n請先把 .xls 另存為 .xlsx，並放到 templates/ 底下。",
            500
        )

    # 清掉舊資料（保留模板樣式）
    _clear_values_keep_style(ws, start_row=2)

    # 模板欄位順序 A~Q（共 17 欄）
    # A 姓名
    # B 身分證字號
    # C 服務日期起（service_start）
    # D 服務日期迄（service_end）
    # E 服務項目
    # F 服務內容
    # G 服務時數_小時
    # H 服務時數_分鐘
    # I 受服務人次
    # J 交通費
    # K 誤餐費
    # L 服務區域
    # M 備註
    # N 匯入動作
    # O 序號（serial_number）
    # P 國外參與服務人次
    # Q 國內參與服務人次

    start_row = 2
    if not records:
        ws.cell(row=start_row, column=1, value="無資料")
    else:
        for i, r in enumerate(records, start=start_row):
            ws.cell(row=i, column=1,  value=r.get("name") or "")
            ws.cell(row=i, column=2,  value=r.get("id_number") or "")
            c3 = ws.cell(row=i, column=3, value=to_roc_yyyMMdd(r.get("service_start")))
            c4 = ws.cell(row=i, column=4, value=to_roc_yyyMMdd(r.get("service_end")))
            c3.number_format = "@"
            c4.number_format = "@"
            ws.cell(row=i, column=5,  value=r.get("service_item") or "")
            ws.cell(row=i, column=6,  value=r.get("service_content") or "")
            h_out, m_out = minutes_to_0_or_30(r.get("service_hours"), r.get("service_minutes"))
            ws.cell(row=i, column=7, value=h_out)
            ws.cell(row=i, column=8, value=m_out)
            ws.cell(row=i, column=9,  value=r.get("served_people_count") or 0)
            ws.cell(row=i, column=10, value=r.get("transport_fee") or 0)
            ws.cell(row=i, column=11, value=r.get("meal_fee") or 0)
            ws.cell(row=i, column=12, value=r.get("service_area") or "")
            ws.cell(row=i, column=13, value=r.get("remarks") or "")
            ws.cell(row=i, column=14, value=r.get("import_action") or "")
            ws.cell(row=i, column=15, value=r.get("serial_number") or "")
            ws.cell(row=i, column=16, value=blank_if_zero(r.get("foreign_service_count")))
            ws.cell(row=i, column=17, value=blank_if_zero(r.get("domestic_service_count")))


    filename = _build_filename(name, nid, date_start, date_end, tag="detailed", ext="xlsx")

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()
    out_bytes, out_ext = convert_xlsx_bytes_to_xls_bytes(xlsx_bytes)

    # 建立安全的 Content-Disposition：提供 ASCII filename 與 UTF-8 percent-encoded filename*
    from urllib.parse import quote
    out_name = filename.replace('.xlsx', '.' + out_ext)
    ascii_name = out_name.encode('ascii', 'ignore').decode('ascii') or 'download.' + out_ext
    disposition = f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(out_name)}'
    response = make_response(out_bytes)
    response.headers["Content-Disposition"] = disposition
    if out_ext == 'xls':
        response.headers["Content-Type"] = "application/vnd.ms-excel"
    else:
        response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response


def export_summary_format(records, name, nid, date_start, date_end):
    """統計格式（按人員統計）- ✅ 使用模板，不自行生成 Workbook"""

    wb, ws = _load_template(TEMPLATE_PATH, SHEET_SUMMARY_NAME)
    if wb is None:
        return make_response(
            f"找不到統計模板檔案：{TEMPLATE_PATH}\n請先準備統計用的 .xlsx 模板並放到 templates/ 底下。",
            500
        )

    # 按姓名分組並統計
    grouped_data = defaultdict(lambda: {
        'id_number': '',
        'service_dates': [],
        'service_items': set(),
        'service_contents': set(),
        'total_hours': 0,
        'total_minutes': 0,
        'total_people': 0,
        'total_transport': 0,
        'total_meal': 0,
        'service_areas': set(),
        'remarks': [],
        'import_actions': set(),
        'serials': [],
        'foreign_count': 0,
        'domestic_count': 0
    })

    for r in records:
        name_val = r.get('name') or ""
        if not name_val:
            continue

        grouped_data[name_val]['id_number'] = r.get('id_number') or ""

        if r.get('service_start'):
            try:
                service_date = r['service_start'].strftime('%Y-%m-%d')
            except Exception:
                service_date = str(r.get('service_start'))
            grouped_data[name_val]['service_dates'].append(service_date)

        if r.get('service_item'):
            grouped_data[name_val]['service_items'].add(str(r['service_item']))
        if r.get('service_content'):
            grouped_data[name_val]['service_contents'].add(str(r['service_content']))

        grouped_data[name_val]['total_hours'] += (r.get('service_hours') or 0)
        grouped_data[name_val]['total_minutes'] += (r.get('service_minutes') or 0)
        grouped_data[name_val]['total_people'] += (r.get('served_people_count') or 0)
        grouped_data[name_val]['total_transport'] += (r.get('transport_fee') or 0)
        grouped_data[name_val]['total_meal'] += (r.get('meal_fee') or 0)

        if r.get('service_area'):
            grouped_data[name_val]['service_areas'].add(str(r['service_area']))

        if r.get('remarks'):
            grouped_data[name_val]['remarks'].append(str(r['remarks']))

        if r.get('import_action'):
            grouped_data[name_val]['import_actions'].add(str(r['import_action']))

        serial_num = r.get('serial_number') or ""
        if serial_num:
            grouped_data[name_val]['serials'].append(str(serial_num))

        grouped_data[name_val]['foreign_count'] += (r.get('foreign_service_count') or 0)
        grouped_data[name_val]['domestic_count'] += (r.get('domestic_service_count') or 0)

    # 處理分鐘進位
    for name_val in grouped_data:
        data = grouped_data[name_val]
        if data['total_minutes'] >= 60:
            extra_hours = data['total_minutes'] // 60
            data['total_hours'] += extra_hours
            data['total_minutes'] = data['total_minutes'] % 60

    # 清掉舊資料（保留模板樣式）
    _clear_values_keep_style(ws, start_row=2)

    start_row = 2
    if not grouped_data:
        ws.cell(row=start_row, column=1, value="無資料")
    else:
        row_i = start_row
        for name_val in sorted(grouped_data.keys()):
            data = grouped_data[name_val]

            dates = sorted(data['service_dates'])
            date_start_str = to_roc_yyyMMdd(dates[0]) if dates else ""
            date_end_str   = to_roc_yyyMMdd(dates[-1]) if dates else ""

            h_out, m_out = minutes_to_0_or_30(data['total_hours'], data['total_minutes'])
            row = [
                name_val,
                data['id_number'],
                date_start_str,
                date_end_str,
                '、'.join(sorted(data['service_items'])) if data['service_items'] else "",
                '、'.join(sorted(data['service_contents'])) if data['service_contents'] else "",
                h_out,
                m_out,
                data['total_people'],
                data['total_transport'],
                data['total_meal'],
                '、'.join(sorted(data['service_areas'])) if data['service_areas'] else "",
                '；'.join(data['remarks']) if data['remarks'] else "",
                '、'.join(sorted(data['import_actions'])) if data['import_actions'] else "",
                '、'.join(data['serials']) if data['serials'] else "",
                blank_if_zero(data['foreign_count']),
                blank_if_zero(data['domestic_count'])

            ]

            for col_i, val in enumerate(row, start=1):
                ws.cell(row=row_i, column=col_i, value=val)

            row_i += 1

    filename = _build_filename(name, nid, date_start, date_end, tag="summary", ext="xlsx")

    output = io.BytesIO()
    wb.save(output)
    xlsx_bytes = output.getvalue()
    out_bytes, out_ext = convert_xlsx_bytes_to_xls_bytes(xlsx_bytes)
    
    from urllib.parse import quote
    out_name = filename.replace('.xlsx', '.' + out_ext)
    ascii_name = out_name.encode('ascii', 'ignore').decode('ascii') or 'download.' + out_ext
    disposition = f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(out_name)}'
    response = make_response(out_bytes)
    response.headers["Content-Disposition"] = disposition
    if out_ext == 'xls':
        response.headers["Content-Type"] = "application/vnd.ms-excel"
    else:
        response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return response

def convert_xlsx_bytes_to_xls_bytes(xlsx_bytes: bytes) -> tuple:
    """
    嘗試使用 LibreOffice headless 將 XLSX bytes 轉成 XLS bytes（Excel 97-2003）。
    若系統沒有安裝 libreoffice 或轉檔失敗，回傳原本的 xlsx bytes 與副檔名 'xlsx'.
    回傳 (bytes, ext)
    """
    try:
        with tempfile.TemporaryDirectory() as tmpdir:
            in_path = os.path.join(tmpdir, "input.xlsx")
            out_path = os.path.join(tmpdir, "input.xls")

            with open(in_path, "wb") as f:
                f.write(xlsx_bytes)

            cmd = [
                "libreoffice",
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                "--convert-to", "xls",
                in_path,
                "--outdir", tmpdir
            ]

            r = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, timeout=60)
            if r.returncode != 0:
                # 若轉檔失敗（例如系統無 libreoffice），回退到 xlsx
                return xlsx_bytes, 'xlsx'

            if not os.path.exists(out_path):
                return xlsx_bytes, 'xlsx'

            with open(out_path, "rb") as f:
                return f.read(), 'xls'
    except Exception:
        # 任何例外都回傳原始 xlsx
        return xlsx_bytes, 'xlsx'

def blank_if_zero(v):
    # None、""、0、"0" 都轉空白
    if v is None or v == "" or v == 0 or v == "0":
        return ""
    return v
